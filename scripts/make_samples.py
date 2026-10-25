"""Generate the sample TC documents in samples/ (needs reportlab: pip install -r requirements-dev.txt).

  tc_221_galaxy_promo.pdf     PDF, product -> gift table, 72-row device table over several pages,
                              repeated page header / footer
  tc_222_mtn_voucher.docx     DOCX with Word heading styles, fixed R500 voucher
  tc_223_choose_gift.md       Markdown, customer chooses a gift
  tc_223_choose_gift_v2.md    same campaign, one more gift option (to try a replacement)

The rules match the validation agent's mock RAG data, so the agent's sample claims give
the same decisions when the agent reads from this service.
"""
from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent.parent / "samples"

COMMON = [
    ("Promotion period", [
        "The promotion is valid for qualifying purchases made between 1 September 2026 and "
        "30 November 2026 (inclusive)."]),
    ("Qualifying devices", [
        "Qualifying devices are the Samsung Galaxy A55, Galaxy S25 and Galaxy S25 Ultra only. "
        "The model codes that qualify are listed in Schedule A."]),
    ("Participating retailers", [
        "Purchases must be made at MTN Stores, Vodacom Shops or Takealot. Purchases from other "
        "retailers do not qualify."]),
    ("Eligibility", ["Open to South African residents with a valid South African ID number."]),
    ("Claim window", ["Claims must be submitted within 30 days of the purchase date."]),
    ("Limits", ["One claim per IMEI/serial number and one claim per ID number per campaign."]),
    ("Proof of purchase", [
        "The invoice must show the invoice number, store name, purchase date and device IMEI or "
        "serial number."]),
]

GENERAL = [
    "9.1 The promoter reserves the right to verify the validity of every claim and to request additional "
    "documents, including a certified copy of the claimant's identity document, before a gift is issued.",
    "9.2 Claims that are incomplete, illegible, fraudulent, submitted through automated means or that do "
    "not comply with these terms and conditions will be rejected without further correspondence.",
    "9.3 Gifts are not transferable and cannot be exchanged for cash. If a gift becomes unavailable for "
    "reasons beyond the promoter's control, the promoter may substitute a gift of equal or greater value.",
    "9.4 Gifts will be delivered within 60 days of a claim being approved. Vouchers are sent by SMS and "
    "email to the contact details provided with the claim; physical gifts are delivered by courier to a "
    "South African street address.",
    "9.5 A device that is returned to the retailer for a refund after a claim is approved forfeits the gift. "
    "Where the gift has already been delivered, the promoter may recover its value from the claimant.",
    "9.6 Employees of the promoter, the participating retailers and their advertising agencies, and their "
    "immediate family members, may not participate in this promotion.",
    "9.7 Personal information is processed in accordance with the Protection of Personal Information Act, "
    "2013, and is used only to administer this promotion unless the claimant consents otherwise.",
    "9.8 The promoter's decision on any claim is final. These terms and conditions are governed by the laws "
    "of the Republic of South Africa.",
]

DEVICES = [
    ("Galaxy A55", ["SM-A556E", "SM-A556E/DS", "SM-A556B"], ["128GB", "256GB"],
     ["Awesome Navy", "Awesome Iceblue", "Awesome Lilac", "Awesome Lemon"]),
    ("Galaxy S25", ["SM-S931B", "SM-S931B/DS"], ["128GB", "256GB", "512GB"],
     ["Icyblue", "Navy", "Mint", "Silver Shadow"]),
    ("Galaxy S25 Ultra", ["SM-S938B", "SM-S938B/DS"], ["256GB", "512GB", "1TB"],
     ["Titanium Silverblue", "Titanium Black", "Titanium Gray", "Titanium Whitesilver"]),
]


def device_rows():
    rows = [["Model", "Model code", "Storage", "Colour"]]
    for model, codes, storages, colours in DEVICES:
        for code in codes:
            for st in storages:
                for col in colours:
                    rows.append([model, code, st, col])
    return rows


# ---------------------------------------------------------------- 221: PDF
def make_pdf(path: Path):
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, leading=20, spaceAfter=8)
    h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=5)
    grid = TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
    ])

    story = [Paragraph("Samsung Galaxy Cashback Promotion 2026 - Terms and Conditions", title)]
    for heading, paras in COMMON:
        story.append(Paragraph(heading, h))
        story += [Paragraph(p, body) for p in paras]

    story.append(Paragraph("Gifts", h))
    story.append(Paragraph("Each valid claim receives the gift listed for the purchased device:", body))
    story.append(Table([["Device", "Gift"],
                        ["Galaxy A55", "R500 Takealot voucher"],
                        ["Galaxy S25", "Samsung Galaxy Buds3"],
                        ["Galaxy S25 Ultra", "Samsung Galaxy Watch7"]], colWidths=[60 * mm, 80 * mm], style=grid))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Gifts are issued only after the claim has been approved.", body))

    story.append(Paragraph("General terms", h))
    story += [Paragraph(p, body) for p in GENERAL]

    story.append(Paragraph("Schedule A - Qualifying device models", h))
    story.append(Paragraph("Only the following model codes qualify for this promotion:", body))
    story.append(Table(device_rows(), repeatRows=1, colWidths=[42 * mm, 36 * mm, 24 * mm, 50 * mm], style=grid))

    def page_marks(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, A4[1] - 12 * mm, "Samsung Galaxy Cashback Promotion 2026 | Terms and Conditions")
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(path), pagesize=A4, topMargin=20 * mm, bottomMargin=18 * mm,
                      leftMargin=20 * mm, rightMargin=20 * mm,
                      title="Samsung Galaxy Cashback Promotion 2026").build(
        story, onFirstPage=page_marks, onLaterPages=page_marks)


# ---------------------------------------------------------------- 222: DOCX
def make_docx(path: Path):
    doc = Document()
    doc.add_heading("MTN Galaxy Voucher Promotion 2026 - Terms and Conditions", level=0)
    for heading, paras in COMMON:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p.replace(" The model codes that qualify are listed in Schedule A.", ""))
    doc.add_heading("Gifts", level=1)
    doc.add_paragraph("Campaign 222 is eligible for an R500 voucher: every successful claim receives an R500 MTN voucher.")
    doc.add_heading("Excluded purchases", level=1)
    for item in ("Refurbished or demo devices.", "Devices bought on a business (B2B) account.",
                 "Purchases made outside South Africa."):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Contact details", level=1)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Channel", "Details"
    for a, b in (("Call centre", "0800 726 786 (Mon-Fri, 08:00-17:00)"),
                 ("Email", "promotions@example.co.za"),
                 ("Website", "www.example.co.za/galaxy-voucher")):
        r = t.add_row().cells
        r[0].text, r[1].text = a, b
    doc.save(str(path))


# ---------------------------------------------------------------- 223: Markdown
def make_md(path: Path, options: list[str]):
    lines = ["# Galaxy Choose-Your-Gift Promotion 2026 - Terms and Conditions", ""]
    for heading, paras in COMMON:
        lines += [f"## {heading}", ""]
        lines += [p.replace(" The model codes that qualify are listed in Schedule A.", "") for p in paras]
        lines.append("")
    lines += ["## Gifts", "",
              "When submitting the claim, the customer may choose one of the following gifts: "
              + ", ".join(options) + ".", "",
              "## Gift values", "",
              "| Gift | Retail value |", "|---|---|"]
    values = {"Samsung Galaxy Buds3": "R2 999", "R750 Takealot voucher": "R750",
              "Samsung 25W Travel Adapter": "R399", "Samsung 45W Travel Adapter": "R699"}
    lines += [f"| {o} | {values[o]} |" for o in options]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_pdf(OUT / "tc_221_galaxy_promo.pdf")
    make_docx(OUT / "tc_222_mtn_voucher.docx")
    base = ["Samsung Galaxy Buds3", "R750 Takealot voucher", "Samsung 25W Travel Adapter"]
    make_md(OUT / "tc_223_choose_gift.md", base)
    make_md(OUT / "tc_223_choose_gift_v2.md", base + ["Samsung 45W Travel Adapter"])
    print("written:", *sorted(p.name for p in OUT.iterdir()), sep="\n  ")
