"""Readers and chunker: headings, tables, clause splitting, PDF / DOCX specifics."""
import io
from pathlib import Path

import pytest
from docx import Document

from app.chunker import chunk_blocks, estimate_tokens
from app.readers import DocumentError, read_document

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def chunks_of(name: str, data: bytes, **kw):
    return chunk_blocks(read_document(name, data), **kw)


# ------------------------------------------------------------------ chunker
def test_table_is_never_cut_and_keeps_its_caption():
    md = b"## Gifts\nEach claim receives the gift below:\n\n| Device | Gift |\n|---|---|\n| A55 | R500 |\n| S25 | Buds3 |\n"
    chunks, _ = chunks_of("t.md", md)
    assert len(chunks) == 1 and chunks[0].kind == "table"
    assert chunks[0].text.splitlines()[0] == "Each claim receives the gift below:"


def test_large_table_split_by_rows_with_header_repeated():
    rows = "\n".join(f"| Model {i} | SM-X{i:03d} | {i * 10} GB |" for i in range(80))
    md = f"## Schedule A\n| Model | Code | Storage |\n|---|---|---|\n{rows}\n".encode()
    chunks, _ = chunks_of("t.md", md, table_max_tokens=300)
    assert len(chunks) > 1 and all(c.kind == "table" for c in chunks)
    body_rows = []
    for c in chunks:
        assert c.tokens <= 300
        assert "| Model | Code | Storage |" in c.text            # header in every part
        assert c.part and c.section == "Schedule A"
        body_rows += [ln for ln in c.text.splitlines() if ln.startswith("| Model ") and "Code" not in ln]
    assert len(body_rows) == 80 and len(set(body_rows)) == 80      # every row once, none lost


def test_long_section_split_at_sentences_within_limit():
    para = " ".join(f"Clause sentence number {i} explains a rule in detail." for i in range(120))
    chunks, _ = chunks_of("t.md", f"## General terms\n{para}\n".encode(), max_tokens=100)
    assert len(chunks) > 3
    for c in chunks:
        assert c.section == "General terms" and c.tokens <= 100
        assert c.text.endswith(".")                                  # never mid-sentence
    joined = " ".join(c.text for c in chunks)
    assert joined.count("Clause sentence") == 120                    # nothing lost, nothing repeated


def test_overlap_repeats_tail_when_enabled():
    para = "\n\n".join(f"Rule {i} applies to every claim made." for i in range(60))
    chunks, _ = chunks_of("t.md", f"## Terms\n{para}\n".encode(), max_tokens=60, overlap_tokens=15)
    assert len(chunks) > 2
    assert chunks[1].text.split("\n\n")[0] in chunks[0].text


def test_sections_never_mix_and_nested_path():
    md = b"""# Title
## Gifts
Gift rule.
### Gift table
| A | B |
|---|---|
| 1 | 2 |
## Limits
One claim per IMEI.
"""
    chunks, title = chunks_of("t.md", md)
    assert title == "Title"
    assert [(c.section, c.kind) for c in chunks] == [
        ("Gifts", "text"), ("Gifts > Gift table", "table"), ("Limits", "text")]


def test_numbered_titles_and_clauses_in_plain_text():
    txt = b"""PROMOTION TERMS

1. Promotion period
1.1 The promotion runs from 1 September 2026 to 30 November 2026.
1.2 Late claims are not accepted.

2. Gifts
2.1 Each valid claim receives an R500 voucher.
"""
    chunks, title = chunks_of("t.txt", txt)
    assert title == "PROMOTION TERMS"
    assert [c.section for c in chunks] == ["1. Promotion period", "2. Gifts"]
    assert "1.1 The promotion" in chunks[0].text and "\n\n1.2 Late" in chunks[0].text


def test_estimate_tokens():
    assert estimate_tokens("abcd" * 100) == 100


# ------------------------------------------------------------------ readers
def test_docx_headings_tables_merged_cells_and_lists():
    doc = Document()
    doc.add_heading("TC Title", level=0)
    doc.add_heading("Gifts", level=1)
    doc.add_paragraph("Each valid claim receives the gift listed:")
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Device", "Gift"
    t.cell(1, 0).text, t.cell(1, 1).text = "Galaxy A55", "R500 voucher"
    t.cell(2, 0).merge(t.cell(2, 1)).text = "All gifts subject to stock"
    doc.add_heading("Exclusions", level=1)
    doc.add_paragraph("Refurbished devices.", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)

    chunks, title = chunks_of("t.docx", buf.getvalue())
    assert title == "TC Title"
    table = chunks[0]
    assert table.section == "Gifts" and table.kind == "table"
    assert "| Galaxy A55 | R500 voucher |" in table.text
    assert "All gifts subject to stock" in table.text and table.text.count("All gifts subject to stock") == 1
    assert chunks[1].section == "Exclusions" and chunks[1].text.startswith("•")


def test_pdf_sample_tables_headers_footers_and_clauses():
    chunks, title = chunks_of("tc.pdf", (SAMPLES / "tc_221_galaxy_promo.pdf").read_bytes())
    assert title.startswith("Samsung Galaxy Cashback Promotion 2026")
    sections = [c.section for c in chunks]
    assert sections[:7] == ["Promotion period", "Qualifying devices", "Participating retailers", "Eligibility",
                            "Claim window", "Limits", "Proof of purchase"]
    text = "\n".join(c.text for c in chunks)
    assert "Page " not in text and "| Terms and Conditions" not in text          # page header/footer gone
    gifts = next(c for c in chunks if c.section == "Gifts" and c.kind == "table")
    assert "| Galaxy S25 Ultra | Samsung Galaxy Watch7 |" in gifts.text

    schedule = [c for c in chunks if c.section.startswith("Schedule A")]
    rows = [ln for c in schedule for ln in c.text.splitlines() if ln.startswith("| Galaxy")]
    assert len(rows) == 72 and len(set(rows)) == 72          # table spread over 3 pages, merged, no repeats
    general = next(c for c in chunks if c.section == "General terms")
    assert general.text.count("\n\n") == 7                    # clauses 9.1 - 9.8 kept apart


def test_text_pdf_only():
    with pytest.raises(DocumentError):
        read_document("x.pdf", b"%PDF-1.4\n%garbage")


def test_image_only_pdf_rejected():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(100, 100, 200, 200, fill=1)        # a drawing, no text: like a scanned page
    c.save()
    with pytest.raises(DocumentError, match="scanned"):
        read_document("scan.pdf", buf.getvalue())


def test_text_encodings():
    blocks = read_document("t.txt", "Gift: R500 voucher – café".encode("cp1252"))
    assert "café" in blocks[0].text
    blocks = read_document("t.txt", "﻿Gift rule".encode("utf-8"))
    assert blocks[0].text == "Gift rule"
